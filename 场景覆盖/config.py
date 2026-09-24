# config.py
"""覆盖率仿真 SSOT — 路径解析 / 列号映射 / 场景解析 / 常量。"""
import sys

from openpyxl.styles import PatternFill

from 公共.路径 import 发现框架路径
from 战斗模拟.config import SIM_WORKBOOK

_mod = sys.modules[__name__]


def resolve_framework_path(base_dir=None, prefer_copy: bool = False) -> str:
    del base_dir, prefer_copy
    found = 发现框架路径()
    if found is not None:
        return str(found)
    return str(SIM_WORKBOOK)


FRAMEWORK_FILE = resolve_framework_path()
# 与战斗模拟共用；启动脚本 -w 会覆写 FRAMEWORK_FILE
SIM_WORKBOOK = FRAMEWORK_FILE  # noqa: F811 — 场景覆盖入口默认

SHEET_OUTPUT = "覆盖率结果"
MATRIX_SHEET = "克制矩阵"
SHEET_SCENE = "场景覆盖"

# 页结构：R1=表名，R2空，R3=块名，R4=字段表头，R5起数据（与框架目录页对齐）。
SCENE_BLOCK_TITLE_ROW = 3
SCENE_HEADER_ROW = 4
SCENE_DATA_ROW = 5

COL: dict = {}

# resolve_columns / resolve_scenes 填充的模块级列表(迁移期兼容, 不用 globals())
ATTR_TYPES: list = []
RACE_TYPES: list = []
SIZE_TYPES: list = []
SCENES: list = []
SCENE_TO_PLAYWAY: dict = {}
N_SCENES: int = 0
SCENE_DATA_N_ROWS: int = 67

# 1 字简称 → 全称(克制矩阵表头归一)
ATTR_NAME_NORMALIZE = {
    "火": "火属性", "水": "水属性", "风": "风属性", "地": "地属性",
    "毒": "毒属性", "圣": "圣属性", "暗": "暗属性", "念": "念属性",
    "无": "无属性", "不死": "不死属性",
}


def _find_col(ws, r4_name: str, block_name: str = None, r3_name: str = None,
              r1_map: dict = None, r3_map: dict = None) -> int:
    """扫块标题行/辅行/字段表头行，找到唯一列号。"""
    matches = []
    for c in range(1, ws.max_column + 1):
        r4 = ws.cell(row=SCENE_HEADER_ROW, column=c).value
        if not r4 or str(r4).strip() != r4_name:
            continue
        if block_name is not None:
            r1_val = (r1_map.get(c) if r1_map is not None
                      else ws.cell(row=SCENE_BLOCK_TITLE_ROW, column=c).value)
            if not r1_val or str(r1_val).strip() != block_name:
                continue
        if r3_name is not None:
            r3_val = (r3_map.get(c) if r3_map is not None
                      else ws.cell(row=3, column=c).value)
            if r3_val is None or str(r3_val).strip() != r3_name:
                continue
        matches.append(c)

    if not matches:
        scope = f"块={block_name!r} " if block_name else ""
        scope += f"辅行={r3_name!r} " if r3_name else ""
        raise KeyError(f"表头未找到: {scope}表头行={r4_name!r}")
    if len(matches) > 1:
        from openpyxl.utils import get_column_letter
        cols_str = ", ".join(f"{get_column_letter(c)}({c})" for c in matches)
        raise KeyError(f"表头不唯一: {r4_name!r} 匹配 {len(matches)} 列: {cols_str}")
    return matches[0]


def _build_inherit_map(ws, row_num: int) -> dict:
    """对块标题行建继承式 dict: None 值继承左侧最近非 None。"""
    result = {}
    last_val = None
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=row_num, column=c).value
        if v is not None and str(v).strip():
            last_val = str(v).strip()
        result[c] = last_val
    return result


# 「场景覆盖」Sheet 以块标题行（R3）为定位锚（禁止写死列号）。
SCENE_R1_BLOCKS = (
    "流派三维",
    "玩法分类",
    "玩法等级分布占比",
    "玩法三维分布",
    "玩法流派偏好",
    "地图遭遇",
    "地图价值",
)

# 地图遭遇块内非三维占比的元数据列（R4）
MAP_META_R4 = {
    "编号", "地图/Boss", "生态区", "玩法类型", "等级",
    "合计价值", "场景数", "平均价值", "备注",
}

# 覆盖率结果：明细数据块自 R25 起（不再写新鲜度横幅行）


def find_r1_title(ws, title: str, *, scan_rows: int = 5,
                  max_col: int | None = None) -> tuple:
    """按块标题（前 scan_rows 行，通常 R3）定位 → (row, col)。找不到 raise KeyError。"""
    want = str(title).strip()
    mc = max_col if max_col is not None else min(int(ws.max_column or 1), MAX_SEARCH_COL + 40)
    last_r = min(int(scan_rows), int(ws.max_row or 1))
    sheet_title = str(getattr(ws, "title", "") or "")
    for r in range(1, last_r + 1):
        for c in range(1, mc + 1):
            v = ws.cell(r, c).value
            if v is None:
                continue
            s = str(v).strip()
            if s == sheet_title and r == 1 and c == 1:
                continue
            if s == want:
                return r, c
    raise KeyError(f"块标题未找到: {want!r}（已扫 R1–R{last_r} × C1–C{mc}）")


def iter_r1_titles(ws, *, scan_rows: int = None, max_col: int | None = None) -> list:
    """从左到右收集块标题行非空标题 [(name, row, col), ...]；同名只留首次。

    默认只扫 SCENE_BLOCK_TITLE_ROW；若传入 scan_rows 则扫 R1..scan_rows（跳过表名）。
    """
    mc = max_col if max_col is not None else min(int(ws.max_column or 1), MAX_SEARCH_COL + 40)
    sheet_title = str(getattr(ws, "title", "") or "")
    if scan_rows is None:
        rows = [SCENE_BLOCK_TITLE_ROW]
    else:
        rows = list(range(1, min(int(scan_rows), int(ws.max_row or 1)) + 1))
    seen = set()
    out = []
    for r in rows:
        for c in range(1, mc + 1):
            v = ws.cell(r, c).value
            if v is None or not str(v).strip():
                continue
            name = str(v).strip()
            if name == sheet_title and r == 1 and c == 1:
                continue
            if name in seen:
                continue
            seen.add(name)
            out.append((name, r, c))
    out.sort(key=lambda t: (t[2], t[1]))
    return out


def block_col_span(ws, block_name: str, *, r1_map: dict = None) -> tuple:
    """命名块的闭区间列 span (start_col, end_col)。

    结束列 = 下一块标题前一列；若无下一块则扫到该块继承范围内最后非空表头格。
    """
    want = str(block_name).strip()
    _, start = find_r1_title(ws, want, scan_rows=5)
    titles = iter_r1_titles(ws)
    next_starts = [c for name, _r, c in titles if c > start]
    if next_starts:
        end = min(next_starts) - 1
    else:
        end = min(int(ws.max_column or start), MAX_SEARCH_COL + 40)
    # 去掉块尾空列（分隔列）
    while end > start:
        h6 = ws.cell(SCENE_HEADER_ROW, end).value
        h4 = ws.cell(4, end).value
        # 0 是合法占比，不能用 `or ""` 当成空
        def _blank(v):
            return v is None or (isinstance(v, str) and not str(v).strip())
        if not _blank(h6) or not _blank(h4):
            break
        end -= 1
    if r1_map is not None:
        # 仍要求继承 R1 名匹配，防止空列被吃进邻块
        while end > start and (r1_map.get(end) or "") != want:
            end -= 1
    return start, end


def _find_header_in_span(ws, start: int, end: int, name: str, rows=(4, 5, 6)) -> tuple[int | None, int | None]:
    """在块列区间内找表头格，兼容地图价值 R4 与流派三维 R6。"""
    want = str(name).strip()
    for r in rows:
        for c in range(start, end + 1):
            v = ws.cell(r, c).value
            if v is not None and str(v).strip() == want:
                return r, c
    return None, None


def _r4_in_span(ws, start: int, end: int, r4_name: str) -> int | None:
    _hr, col = _find_header_in_span(ws, start, end, r4_name)
    return col


def _classify_map_r4(name: str) -> str:
    """地图遭遇 R4 → attr / race / size / meta。

    注意：ATTR_NAME_NORMALIZE 短键（如「不死」→「不死属性」）是克制矩阵用的，
    地图遭遇里「不死」是种族，「不死属性」才是属性——不能用短键当属性判定。
    """
    s = str(name).strip()
    if s in MAP_META_R4:
        return "meta"
    if s.endswith("体型") or s in ("小体型", "中体型", "大体型", "小", "中", "大"):
        return "size"
    if s.endswith("属性") or s in set(ATTR_NAME_NORMALIZE.values()):
        return "attr"
    return "race"


def resolve_columns(ws) -> dict:
    """按块标题行 + 字段表头行自动发现列；填充 COL 与 ATTR/RACE/SIZE_TYPES。"""
    r1_map = _build_inherit_map(ws, SCENE_BLOCK_TITLE_ROW)
    COL.clear()

    # 1) 登记全部 R1 块 span（列号可变，标题不可变）
    missing_blocks = []
    for bname in SCENE_R1_BLOCKS:
        try:
            s, e = block_col_span(ws, bname, r1_map=r1_map)
        except KeyError:
            missing_blocks.append(bname)
            continue
        COL[f"BLOCK_{bname}_START"] = s
        COL[f"BLOCK_{bname}_END"] = e
        COL[f"块_{bname}_START"] = s
        COL[f"块_{bname}_END"] = e

    if missing_blocks:
        qprint(f"{PRINT_PREFIX_WARN} 块标题未找到: {', '.join(missing_blocks)}")

    # 2) {块名}_{R4} → 列号（仅块内，避免跨块同名 R4）
    for c in range(1, ws.max_column + 1):
        block = r1_map.get(c)
        for r in (4, SCENE_HEADER_ROW):
            col_name = ws.cell(r, c).value
            if col_name is None or (isinstance(col_name, str) and not col_name.strip()):
                continue
            if not isinstance(col_name, str):
                continue
            key = f"{block}_{col_name.strip()}" if block else col_name.strip()
            COL.setdefault(key, c)

    def _span(bname: str):
        s = COL.get(f"BLOCK_{bname}_START")
        e = COL.get(f"BLOCK_{bname}_END")
        if s is None or e is None:
            return None
        return int(s), int(e)

    # 3) 地图遭遇：按 R4 名分类属性/种族/体型，不再写死 +10/+20
    sp_map = _span("地图遭遇")
    if sp_map:
        s, e = sp_map
        attrs, races, sizes = [], [], []
        attr_cols, race_cols, size_cols = [], [], []
        for c in range(s, e + 1):
            v = ws.cell(4, c).value
            if v is None or (isinstance(v, str) and not str(v).strip()) or not isinstance(v, str):
                v = ws.cell(SCENE_HEADER_ROW, c).value
            if v is None or not isinstance(v, str) or not str(v).strip():
                continue
            name = str(v).strip()
            kind = _classify_map_r4(name)
            if kind == "attr":
                attrs.append(name)
                attr_cols.append(c)
            elif kind == "race":
                races.append(name)
                race_cols.append(c)
            elif kind == "size":
                sizes.append(name)
                size_cols.append(c)
        if attr_cols:
            COL["MAP_ATTR_START"] = attr_cols[0]
            COL["MAP_ATTR_COUNT"] = len(attr_cols)
        if race_cols:
            COL["MAP_RACE_START"] = race_cols[0]
            COL["MAP_RACE_COUNT"] = len(race_cols)
        if size_cols:
            COL["MAP_SIZE_START"] = size_cols[0]
            COL["MAP_SIZE_COUNT"] = len(size_cols)
        if attrs:
            ATTR_TYPES.clear()
            ATTR_TYPES.extend(attrs)
        if races:
            RACE_TYPES.clear()
            RACE_TYPES.extend(races)
        if sizes:
            SIZE_TYPES.clear()
            SIZE_TYPES.extend(sizes)

    # 4) 玩法三维分布：PVE%/PVP%/玩法 必须落在该 R1 块内
    sp_pw3 = _span("玩法三维分布")
    if sp_pw3:
        s, e = sp_pw3
        pve_c = _r4_in_span(ws, s, e, "PVE%") or COL.get("玩法三维分布_PVE%")
        pvp_c = _r4_in_span(ws, s, e, "PVP%") or COL.get("玩法三维分布_PVP%")
        pw_c = _r4_in_span(ws, s, e, "玩法") or COL.get("玩法三维分布_玩法")
        if pve_c is None or pvp_c is None:
            raise KeyError(
                "未找到玩法三维分布的 PVE%/PVP% 列 "
                f"(PVE%={pve_c}, PVP%={pvp_c}); 请检查 R1「玩法三维分布」与 R4 表头"
            )
        COL["BE_PVE_PCT_COL"] = pve_c
        COL["BE_PVP_PCT_COL"] = pvp_c
        COL["玩法三维分布_PVE%"] = pve_c
        COL["玩法三维分布_PVP%"] = pvp_c
        if pw_c:
            COL["玩法三维分布_玩法"] = pw_c

    # 5) 玩法流派偏好（列范围按块 span，不再假设「玩法」右侧连续 50 列）
    sp_pref = _span("玩法流派偏好")
    if sp_pref:
        s, e = sp_pref
        col_play = _r4_in_span(ws, s, e, "玩法") or COL.get("玩法流派偏好_玩法")
        if col_play:
            build_cols = []
            for c in range(s, e + 1):
                if c == col_play:
                    continue
                v = ws.cell(SCENE_HEADER_ROW, c).value
                if v is None or not str(v).strip():
                    continue
                if str(v).strip() == "玩法":
                    continue
                build_cols.append(c)
            n = len(build_cols)
            if n > 0:
                COL["BE_PLAYWAY_COL"] = col_play
                COL["BE_BUILD_COL_START"] = build_cols[0]
                COL["BE_BUILD_COL_END"] = build_cols[-1] + 1
                COL["N_BE_BUILDS"] = n
                n_pw = 0
                empty = 0
                for r in range(BE_DATA_ROW_START, ws.max_row + 1):
                    v = ws.cell(row=r, column=col_play).value
                    if v is None or not str(v).strip():
                        empty += 1
                        if empty >= 2:
                            break
                        continue
                    empty = 0
                    if str(v).strip() == "玩法":
                        continue
                    n_pw += 1
                COL["N_BE_PLAYWAYS"] = max(n_pw, BE_DATA_ROW_END - BE_DATA_ROW_START + 1)

    # 6) 玩法等级分布占比 / 玩法分类：只登记关键列（供后续 loader 用块名而非列号）
    sp_lv = _span("玩法等级分布占比")
    if sp_lv:
        s, e = sp_lv
        for r4 in ("等级", "价值占比", "时长占比", "PVE价值占比", "PVP价值占比"):
            c = _r4_in_span(ws, s, e, r4)
            if c:
                COL[f"玩法等级分布占比_{r4}"] = c

    sp_cls = _span("玩法分类")
    if sp_cls:
        s, e = sp_cls
        for r4 in ("玩法", "定位", "类型", "开放等级"):
            c = _r4_in_span(ws, s, e, r4)
            if c:
                COL[f"玩法分类_{r4}"] = c

    sp_spec = _span("流派三维")
    if sp_spec:
        s, e = sp_spec
        cat = _r4_in_span(ws, s, e, "大类")
        sub = _r4_in_span(ws, s, e, "小类")
        if cat:
            COL["流派三维_大类"] = cat
        if sub:
            COL["流派三维_小类"] = sub
        build_cols = []
        for c in range(s, e + 1):
            if c in (cat, sub):
                continue
            v = ws.cell(SCENE_HEADER_ROW, c).value
            if v is None or not str(v).strip():
                continue
            build_cols.append(c)
        if build_cols:
            COL["SPEC3_BUILD_COL_START"] = build_cols[0]
            COL["SPEC3_BUILD_COL_END"] = build_cols[-1] + 1
            COL["N_SPEC3_BUILDS"] = len(build_cols)

    # 7) 地图价值数据尾行
    col_lv = COL.get("地图价值_等级")
    if col_lv is None:
        sp_val = _span("地图价值")
        if sp_val:
            col_lv = _r4_in_span(ws, sp_val[0], sp_val[1], "等级")
    if col_lv:
        last_row = SCENE_DATA_START
        for r in range(SCENE_DATA_START, ws.max_row + 1):
            v = ws.cell(row=r, column=col_lv).value
            if v is not None and str(v).strip():
                last_row = r
        COL["SCENE_DATA_END"] = last_row
        COL["地图价值_等级"] = col_lv

    found = [b for b in SCENE_R1_BLOCKS if f"BLOCK_{b}_START" in COL]
    qprint(
        f"{PRINT_PREFIX_RESOLVE} 块定位 {len(found)}/{len(SCENE_R1_BLOCKS)} "
        f"({', '.join(found)}); auto-discovered {len(COL)} columns from 块标题+表头"
    )
    return COL


LEVELS_5STEP = [10, 15, 20, 25, 30, 35, 40, 45, 50, 55, 60]


def resolve_scenes(ws):
    """从「地图价值/场景类型」列解析场景名 SSOT, 填充 SCENES / SCENE_TO_PLAYWAY。"""
    r1_map = _build_inherit_map(ws, SCENE_BLOCK_TITLE_ROW)
    scene_col = None
    header_row = None
    try:
        s, e = block_col_span(ws, "地图价值", r1_map=r1_map)
        header_row, scene_col = _find_header_in_span(ws, s, e, "场景类型")
    except KeyError:
        scene_col = None
    if scene_col is None:
        for r in range(3, 8):
            for c in range(1, (ws.max_column or 1) + 1):
                if str(ws.cell(r, c).value or "").strip() == "场景类型":
                    header_row, scene_col = r, c
                    break
            if scene_col is not None:
                break
    if scene_col is None:
        raise ConfigDriftError(
            "场景覆盖 Sheet 未找到「地图价值/场景类型」列 — "
            "场景名 SSOT 应写入该列, 不允许写死列号兜底"
        )
    data_start = int(header_row) + 1
    raw = []
    for r in range(data_start, (ws.max_row or data_start) + 1):
        v = ws.cell(r, scene_col).value
        if v is None or not str(v).strip():
            continue
        s = str(v).strip()
        if "合计" in s:
            continue
        raw.append(s)
    seen = set()
    scenes = []
    dropped = []
    for s in raw:
        if s in seen:
            dropped.append(s)
            continue
        seen.add(s)
        scenes.append(s)
    if dropped:
        qprint(f"{PRINT_PREFIX_RESOLVE} 场景类型列去重: {len(raw)} 行 → {len(scenes)} 唯一场景 "
               f"(丢 {len(dropped)} 个等级×场景重复)")
    playway_col = None
    for c in range(1, ws.max_column + 1):
        r4 = ws.cell(SCENE_HEADER_ROW, c).value
        if r4 and ("流派偏好名" in str(r4) or "playway" in str(r4).lower()):
            playway_col = c
            break
    scene_to_playway = {scene: scene for scene in scenes}
    if playway_col is not None:
        for i, scene in enumerate(scenes):
            v = ws.cell(SCENE_DATA_START + i, playway_col).value
            if v:
                scene_to_playway[scene] = str(v).strip()

    SCENES.clear()
    SCENES.extend(scenes)
    SCENE_TO_PLAYWAY.clear()
    SCENE_TO_PLAYWAY.update(scene_to_playway)
    _mod.N_SCENES = len(scenes)
    _mod.SCENE_DATA_N_ROWS = len(scenes)
    COL["SCENE_DATA_N_ROWS"] = len(scenes)
    COL["N_SCENES_FROM_EXCEL"] = len(scenes)
    mode = "正常" if playway_col else "退化(值=场景名)"
    qprint(f"{PRINT_PREFIX_RESOLVE} 场景解析完成: N={len(scenes)} (地图价值/场景类型 col{scene_col}), "
           f"SCENE_TO_PLAYWAY {mode}")
    return scenes, scene_to_playway


# ---- 流派面板行号 ----
ELEM_DMG_START_ROW = 7
ELEM_DMG_END_ROW = 16
RACE_START_ROW = 35
RACE_END_ROW = 44
SIZE_START_ROW = 45
SIZE_END_ROW = 47
WEAPON_START_ROW = 17
WEAPON_END_ROW = 34

LEVEL_SUMMARY_DATA_START = 5
LEVEL_SUMMARY_DATA_END = 15

SCENE_DATA_TOP_ROW = 5
MAP_DATA_START = 5
SCENE_DATA_START = 5
SCENE_DATA_END = None

PROGRAM_EPS = 1e-9
USER_EPS = 0.01
PVP_BUILD_EPS = 0.05
PVP_WEAPON_EPS = 0.1

BE_DATA_ROW_START = 5
BE_DATA_ROW_END = 27
N_BE_PLAYWAYS = BE_DATA_ROW_END - BE_DATA_ROW_START + 1

NUM_FMT = '0.000000'
INT_FMT = '0'
PCT_FMT = '0.00%'
MAX_SEARCH_COL = 119

BLOCK_GROUP_ROWS = [27, 51, 66, 81]

# ---------------------------------------------------------------------------
# 「覆盖率结果」展示术语（2026-08-04 统一）
#
#   输出乘区 = 克制矩阵期望伤害倍率（打别人，越高越好）
#   承伤乘区 = 克制矩阵期望挨打倍率（被打，越高越疼）
#   平衡/克制/抗性指数 = 维内均值归一后的相对指数（1.0=维内平均）
#   偏离均值% = (指数−1)×100
#   相对均值差 = 均值承伤 − 本项承伤（正=比平均少挨，乘区绝对值）
#   场景占比 = 环境出现权重（种族无克制矩阵时用，非战斗强弱）
#   判定 = PASS/WARN/FAIL（只看指数是否落在阈值带）
# ---------------------------------------------------------------------------

# 五维排名 PASS/WARN/FAIL 带必须来自公式参数；框架簿暂无这些行则只写指数、不发明判定带。
RANK_REL_THRESHOLDS: dict = {}

# β 双源对照容差：场景行推导 β̂ vs 等级曲线 β（excel 覆盖口径）允许的最大漂移。
# 超过即视为两份输入数据脱节，体检 A 组 FAIL（2026-07-30 登记）。
BETA_DRIFT_TOLERANCE = 0.10

# 面板「武器类型」与克制矩阵细名已对齐；桥接为 1:1（兼容旧粗名别名）
WEAPON_NAME_BRIDGE: dict = {
    "空手": ["空手"],
    "大剑": ["大剑"],
    "剑盾": ["剑盾"],
    "长枪": ["长枪"],
    "弓": ["弓"],
    "鞭子": ["鞭子"],
    "乐器": ["乐器"],
    "盾杖": ["盾杖"],
    "杖书": ["杖书"],
    "法杖": ["法杖"],
    "短剑": ["短剑"],
    "拳套": ["拳套"],
    "巨斧": ["巨斧"],
    # 旧粗名兼容
    "单手剑": ["剑盾"],
    "双手剑": ["大剑"],
    "单手矛": ["长枪"],
    "双手矛": ["长枪"],
    "单手斧": ["巨斧"],
    "双手斧": ["巨斧"],
    "钝器": ["拳套"],
    "单手杖": ["盾杖"],
    "双手杖": ["法杖"],
    "魔法书": ["杖书"],
    "法器": ["法杖"],
    "短弓": ["弓"],
    "长弓": ["弓"],
    "匕首": ["短剑"],
    "拳刃": ["拳套"],
}

WEAPON_MATRIX_TYPES: tuple = (
    "空手", "大剑", "剑盾", "长枪", "弓", "鞭子", "乐器",
    "盾杖", "杖书", "法杖", "短剑", "拳套", "巨斧",
)

# (block_id, 块标题, 起始列, 条数)
# n_data 为看板行数；流派条数在 resolve/load 后按表头动态改写
RANK_BLOCK_LAYOUT = [
    ("build",   "流派平衡", 1,  11),
    ("elem",    "属性克制", 8,  10),
    ("race",    "种族场景占比", 15, 10),
    ("size",    "体型承伤", 21, 3),
    ("weapon",  "武器克制", 27, 13),
]

RANK_HEADERS_ATK_DEF = ["名称", "输出乘区", "承伤乘区", "平衡指数", "偏离均值%", "判定"]
RANK_HEADERS_WEAPON = ["武器", "输出乘区", "克制指数", "偏离均值%", "判定"]
RANK_HEADERS_SIZE = ["体型", "承伤乘区", "相对均值差", "抗性指数", "判定"]
RANK_HEADERS_RACE_SHARE = ["种族", "场景占比", "相对均值", "偏离均值%", "判定"]
RANK_HEADERS_RACE_MATRIX = ["种族", "输出乘区", "承伤乘区", "平衡指数", "偏离均值%", "判定"]

RANK_BLOCK_HEADERS = {
    "build": ["流派", "输出乘区", "承伤乘区", "平衡指数", "偏离均值%", "判定"],
    "elem": ["属性", "输出乘区", "承伤乘区", "克制指数", "偏离均值%", "判定"],
    "race": RANK_HEADERS_RACE_SHARE,
    "size": RANK_HEADERS_SIZE,
    "weapon": RANK_HEADERS_WEAPON,
}

# 兼容旧名
RANK_RACE_HEADERS = RANK_HEADERS_RACE_SHARE

# 数据块标题模板（{env}=综合/PVE/PVP）
BLOCK_TITLE_QUAD = "场景属性乘区·攻守"
BLOCK_TITLE_BFI_CROSS = "跨等级·流派平衡指数"
BLOCK_TITLE_BFI_BY_LV = "{env}·流派平衡指数(分等级)"
BLOCK_TITLE_SCENE_DIST = "{env}·{dim}场景占比"
BLOCK_TITLE_ATTR_ATK = "{env}·属性输出乘区"
BLOCK_TITLE_ATTR_DEF = "{env}·属性承伤乘区"
BLOCK_TITLE_WEAPON_ATK = "{env}·武器对体型输出乘区"
BLOCK_TITLE_SIZE_HIT = "{env}·体型承伤乘区"

BLOCK_HDR_QUAD = ["等级", "PVE输出", "PVE承伤", "PVP输出", "PVP承伤", "综合输出", "综合承伤"]
BLOCK_HDR_BFI = ["流派", "输出乘区", "承伤乘区", "平衡指数"]

# 结果表备注（人话口径）
RESULT_NOTES_BASE = (
    "口径=克制矩阵期望乘区，不含技能循环/蓝耗/控制/命中暴击/面板攻速",
    "综合=β·PVE+(1−β)·PVP；β=等级PVE价值占比(AF列)；等级权重W_L=AD列",
    "输出乘区=打环境的期望倍率；承伤乘区=被环境打的期望倍率(越高越疼)",
    "指数列=维内均值归一(1.0=平均)，供判定；偏离均值%=(指数−1)×100",
    "体型相对均值差=均值承伤−本项承伤(正=比平均少挨，乘区绝对值)",
)

CHART_ANCHOR_ROWS = [97, 97, 97, 97, 97]

FILL_PASS = PatternFill("solid", fgColor="C6EFCE")
FILL_WARN = PatternFill("solid", fgColor="FFEB9C")
FILL_FAIL = PatternFill("solid", fgColor="FFC7CE")

# 「覆盖率结果」表样式
FILL_TITLE = PatternFill("solid", fgColor="0F2744")
FILL_HEADER = PatternFill("solid", fgColor="1B3A5F")
FILL_LABEL = PatternFill("solid", fgColor="F2F2F2")
FILL_SHARE_1 = PatternFill("solid", fgColor="DEEBF7")
FILL_SHARE_2 = PatternFill("solid", fgColor="BDD7EE")
FILL_SHARE_3 = PatternFill("solid", fgColor="9BC2E6")
FILL_SHARE_4 = PatternFill("solid", fgColor="5B9BD5")

# 场景占比热力分档 (绝对占比)
SHARE_FILL_BANDS = (
    (0.05, FILL_SHARE_1),
    (0.10, FILL_SHARE_2),
    (0.20, FILL_SHARE_3),
    (1.01, FILL_SHARE_4),
)

STANDARD_BUILD_NAME = "标准木桩"

QUIET = False


def qprint(*args, **kwargs):
    """QUIET 时 skip,非 QUIET 时等同 print()。警告请用原 print() 强制播报。"""
    if not QUIET:
        print(*args, **kwargs)


class ConfigDriftError(RuntimeError):
    """Excel SSOT 缺数据时 raise — fail-closed。"""


PRINT_PREFIX_READ = "[读取]"
PRINT_PREFIX_WARN = "[警告]"
PRINT_PREFIX_CHAIN_OK = "[链路-OK]"
PRINT_PREFIX_CHAIN_BLOCK = "[链路-阻断]"
PRINT_PREFIX_CHAIN_WARN = "[链路-警告]"
PRINT_PREFIX_CHAIN_DONE = "[链路-完备]"
PRINT_PREFIX_FALLBACK = "[已禁兜底]"
PRINT_PREFIX_AGGREGATE = "[聚合]"
PRINT_PREFIX_QUAD = "[四象限]"
PRINT_PREFIX_BFI = "[BFI]"
PRINT_PREFIX_VALIDATE = "[校验]"
PRINT_PREFIX_CHAIN_BREAK = "[X]"
PRINT_PREFIX_MATRIX_MERGE = "[武器映射]"
PRINT_PREFIX_OUTPUT = "[输出]"
PRINT_PREFIX_RESOLVE = "[列解析]"

VALIDATE_DIST_EPS = 0.001
KEYWORD_SEARCH_MAX_ROWS = 200
