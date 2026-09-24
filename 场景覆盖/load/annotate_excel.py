# annotate_excel.py — 仅加批注, 不改业务单元格数值
from __future__ import annotations

import openpyxl
from openpyxl.comments import Comment

from 场景覆盖 import config as cfg


_NOTE = (
    "【覆盖率工具·玩法三维分布】\n"
    "LV10~LV60 = 各玩法在该等级的活跃/价值份额（列内再归一）。\n"
    "PVE%/PVP% = 该玩法内 PVE 与 PVP 拆分。\n"
    "右侧「玩法流派偏好」= 该玩法下各流派占比。\n"
    "仿真: 每等级玩家 meta = 份额×PVE%(或PVP%)×流派分布；"
    "再与场景怪物三维、等级价值占比合成终身环境。\n"
    "输出表「覆盖率结果」用语: 输出/承伤乘区=矩阵期望；"
    "指数=维内均值归一(1.0=平均)；偏离均值%=(指数−1)×100。\n"
    "请保持左侧玩法名与右侧 BE 玩法名一致。"
)

# 关键输入块的轻量批注（定位靠 R1 标题，不写死列号）
_BLOCK_NOTES = {
    "流派三维": (
        "【覆盖率工具·流派三维】\n"
        "大类/小类 + 各构筑列权重；与右侧玩法流派偏好目录应对齐。\n"
        "构筑列数从表头动态读取（当前与标准模型 11 构筑对齐）。"
    ),
    "玩法分类": (
        "【覆盖率工具·玩法分类】\n"
        "玩法 / 定位 / 类型 / 开放等级；玩法名需与「玩法三维分布」「玩法流派偏好」一致。"
    ),
    "玩法等级分布占比": (
        "【覆盖率工具·玩法等级分布占比】\n"
        "等级段价值/时长及 PVE·PVP 拆分；为综合 β 与等级权重输入。"
    ),
    "玩法三维分布": _NOTE,
}


def annotate_playway_usage(xlsx_path: str | None = None) -> None:
    path = xlsx_path or cfg.FRAMEWORK_FILE
    wb = openpyxl.load_workbook(path)
    if cfg.SHEET_SCENE not in wb.sheetnames:
        wb.close()
        return
    ws = wb[cfg.SHEET_SCENE]

    annotated = []
    for title, note in _BLOCK_NOTES.items():
        try:
            r, c = cfg.find_r1_title(ws, title, scan_rows=5)
        except KeyError:
            continue
        cell = ws.cell(r, c)
        # 只写批注, 绝不改 value (避免破坏合并区/公式锚点)
        cell.comment = Comment(note, "场景覆盖")
        annotated.append(title)

    if not annotated:
        wb.close()
        return

    wb.save(path)
    wb.close()
    cfg.qprint(
        f"{cfg.PRINT_PREFIX_OUTPUT} 已为 R1 块加批注: {', '.join(annotated)}"
    )
