# scenes.py — 场景覆盖 / 地图遭遇 / 等级权重读取与聚合
from collections import defaultdict
from typing import Dict, List, Tuple

import numpy as np
from openpyxl.utils import get_column_letter

from 场景覆盖 import config as cfg
from 场景覆盖.config import ConfigDriftError


def _col_tag(key: str) -> str:
    """按解析后的列号生成 'AD(30)' 式标签；未解析到时退回键名。
    报错文案禁止写死列号（曾写 C31/C33/C34，与实际 AD/AF/AG 漂移，2026-07-30 修）。"""
    c = cfg.COL.get(key)
    if isinstance(c, int):
        return f"{get_column_letter(c)}({c})"
    return f"{key}列"


def _safe_float(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    if isinstance(value, str):
        s = value.strip()
        if s == "":
            raise ConfigDriftError(
                f"_safe_float 遇到空字符串 (业务真值应写入 xlsx, 不允许 cfg fallback)"
            )
        if s.startswith("#") or s.upper() in {"N/A", "NA", "NULL", "NONE", "ERROR", "ERR"}:
            raise ConfigDriftError(
                f"_safe_float 遇到 Excel 异常值 '{s}' (公式返回 #N/A / #REF! / #VALUE!, "
                f"业务真值应写入公式依赖的源数据)"
            )
        has_percent = s.endswith('%')
        s = s.rstrip('%')
        try:
            v = float(s)
        except ValueError:
            raise ConfigDriftError(
                f"_safe_float 无法解析字符串 '{s}' (业务真值应写入 xlsx, 不允许 cfg fallback)"
            )
        return v / 100.0 if has_percent else v
    raise ConfigDriftError(
        f"_safe_float 收到非数值类型 {type(value).__name__}={value!r} "
        f"(业务真值应写入 xlsx, 不允许 cfg fallback)"
    )


def _find_sheet(wb, keyword: str) -> str:
    """在 workbook 中按关键词查找 Sheet 名。"""
    for name in wb.sheetnames:
        if keyword in name:
            return name
    raise KeyError(f"找不到包含'{keyword}'的Sheet")


def _find_keyword_row(ws, keyword: str, start_row: int = 1,
                      max_rows: int = None, col: int = None) -> int:
    """搜索关键词，返回所在行号，未找到返回 -1。col=None 时跨所有列搜索。"""
    if max_rows is None:
        max_rows = cfg.KEYWORD_SEARCH_MAX_ROWS
    if col is not None:
        for r in range(start_row, min(ws.max_row + 1, start_row + max_rows)):
            val = ws.cell(row=r, column=col).value
            if val and keyword in str(val).strip():
                return r
    else:
        for r in range(start_row, min(ws.max_row + 1, start_row + max_rows)):
            for c in range(1, min(ws.max_column + 1, cfg.MAX_SEARCH_COL)):
                val = ws.cell(row=r, column=c).value
                if val and keyword in str(val).strip():
                    return r
    return -1


def read_level_pve_pvp_ratios(wb) -> Dict:
    """读 Excel PVE/PVP 价值占比 → {lv: {pve_weight, pvp_weight}} (sum=1)。"""
    sheet_name = _find_sheet(wb, cfg.SHEET_SCENE)
    ws = wb[sheet_name]

    raw = {}
    for r in range(cfg.LEVEL_SUMMARY_DATA_START, cfg.LEVEL_SUMMARY_DATA_END + 1):
        lv_val = ws.cell(row=r, column=cfg.COL["玩法等级分布占比_等级"]).value
        if lv_val is None:
            continue
        lv_str = str(lv_val).strip()
        if lv_str.upper().startswith("LV"):
            lv_str = lv_str[2:]
        try:
            lv = int(float(lv_str))
        except (ValueError, TypeError):
            continue
        pve = _safe_float(ws.cell(row=r, column=cfg.COL["玩法等级分布占比_PVE价值占比"]).value)
        pvp = _safe_float(ws.cell(row=r, column=cfg.COL["玩法等级分布占比_PVP价值占比"]).value)
        if pve > 0 or pvp > 0:
            total = pve + pvp
            if total > 0:
                raw[lv] = {"pve_weight": pve / total, "pvp_weight": pvp / total}

    if not raw:
        raise ConfigDriftError(
            f"Excel {_col_tag('玩法等级分布占比_PVE价值占比')}/{_col_tag('玩法等级分布占比_PVP价值占比')} "
            "列 (PVE/PVP 价值占比) 为空 — "
            f"业务真值应写入 场景覆盖 Sheet 对应列 R5-R15, 不允许 cfg fallback"
        )

    cfg.qprint(f"{cfg.PRINT_PREFIX_READ} 等级 PVE/PVP 比例 "
               f"({_col_tag('玩法等级分布占比_PVE价值占比')}/{_col_tag('玩法等级分布占比_PVP价值占比')}, "
               f"{len(raw)} 等级)")
    return raw


def read_level_value_ratios(wb) -> Dict[int, float]:
    """等级段价值占比 (Excel SSOT), 归一化后 sum=1.0。"""
    sheet_name = _find_sheet(wb, cfg.SHEET_SCENE)
    ws = wb[sheet_name]

    raw = {}
    for r in range(cfg.LEVEL_SUMMARY_DATA_START, cfg.LEVEL_SUMMARY_DATA_END + 1):
        lv_val = ws.cell(row=r, column=cfg.COL["玩法等级分布占比_等级"]).value
        if lv_val is None:
            continue
        lv_str = str(lv_val).strip()
        if lv_str.upper().startswith("LV"):
            lv_str = lv_str[2:]
        try:
            lv = int(float(lv_str))
        except (ValueError, TypeError):
            continue
        ratio = _safe_float(ws.cell(row=r, column=cfg.COL["玩法等级分布占比_价值占比"]).value)
        if ratio > 0:
            raw[lv] = ratio

    if not raw:
        raise ConfigDriftError(
            f"Excel {_col_tag('玩法等级分布占比_价值占比')} 列 (等级段价值占比) 为空 — "
            "业务真值应写入 场景覆盖 Sheet 对应列 R5-R15, 不允许 cfg fallback"
        )

    total = sum(raw.values())
    if total <= 0:
        raise ConfigDriftError(
            f"Excel {_col_tag('玩法等级分布占比_价值占比')} 列 (等级段价值占比) 存在但总和为 0 — "
            "业务真值应写入 场景覆盖 Sheet 对应列 R5-R15, 不允许均匀兜底"
        )
    result: Dict[int, float] = {}
    sorted_lvs = sorted(raw.keys())
    partial = 0.0
    for lv in sorted_lvs[:-1]:
        w = raw[lv] / total
        result[lv] = w
        partial += w
    result[sorted_lvs[-1]] = 1.0 - partial

    cfg.qprint(f"{cfg.PRINT_PREFIX_READ} 等级段价值占比 ({_col_tag('玩法等级分布占比_价值占比')}, "
               f"Lv60={result.get(60, 0)*100:.1f}%, sum={sum(result.values()):.6f})")
    return result


def read_map_encounters(wb):
    """读取「地图遭遇」→ {地图名称: {attr, race, size}}。"""
    sheet_name = _find_sheet(wb, cfg.SHEET_SCENE)
    ws = wb[sheet_name]

    title_row = _find_keyword_row(ws, "地图遭遇")
    if title_row < 0:
        title_row = _find_keyword_row(ws, "怪物分布")
    if title_row < 0:
        title_row = _find_keyword_row(ws, "地图")

    if title_row < 0:
        cfg.qprint(f"{cfg.PRINT_PREFIX_READ} 地图遭遇: 未找到标题行,尝试使用默认位置")
    else:
        cfg.qprint(f"{cfg.PRINT_PREFIX_READ} 地图遭遇标题定位完成")

    data_start = cfg.MAP_DATA_START
    result = {}
    empty_count = 0
    for r in range(data_start, ws.max_row + 1):
        name = ws.cell(row=r, column=cfg.COL["地图遭遇_地图/Boss"]).value
        if not name:
            empty_count += 1
            if empty_count >= 3:
                break
            continue
        empty_count = 0
        name = str(name).strip()

        playtype = ws.cell(row=r, column=cfg.COL["地图遭遇_玩法类型"]).value
        if not playtype:
            continue
        playtype = str(playtype).strip()
        if playtype not in cfg.SCENE_TO_PLAYWAY:
            for k in cfg.SCENE_TO_PLAYWAY:
                if k.startswith(playtype) or playtype.startswith(k):
                    playtype = k
                    break

        attr = np.array([_safe_float(ws.cell(row=r, column=c).value)
                        for c in range(cfg.COL["MAP_ATTR_START"],
                                       cfg.COL["MAP_ATTR_START"] + 10)])
        race = np.array([_safe_float(ws.cell(row=r, column=c).value)
                         for c in range(cfg.COL["MAP_RACE_START"],
                                        cfg.COL["MAP_RACE_START"] + cfg.COL["MAP_RACE_COUNT"])])
        if len(race) < 10:
            race = np.concatenate([race, np.zeros(10 - len(race))])
        size = np.array([_safe_float(ws.cell(row=r, column=c).value)
                        for c in range(cfg.COL["MAP_SIZE_START"],
                                       cfg.COL["MAP_SIZE_START"] + 3)])
        result[name] = {"attr": attr, "race": race, "size": size}

    cfg.qprint(f"{cfg.PRINT_PREFIX_READ} 地图遭遇 (从 Layer 2)")
    return result


def read_scene_coverage(wb):
    """从场景覆盖 Sheet 读场景价值分配。

    返回: (levels, scene_config, attr_types, size_types, race_types)
    """
    sheet_name = _find_sheet(wb, cfg.SHEET_SCENE)
    ws = wb[sheet_name]

    attr_types = list(cfg.ATTR_TYPES)
    if not attr_types:
        raise ConfigDriftError(
            "ATTR_TYPES 为空 — resolve_columns 未从「地图遭遇」表头解析到属性列"
        )

    size_types = list(cfg.SIZE_TYPES) if cfg.SIZE_TYPES else []
    if not size_types:
        for r in range(cfg.SIZE_START_ROW, cfg.SIZE_END_ROW + 1):
            v = ws.cell(row=r, column=2).value
            if v:
                s = str(v).strip()
                if "小" in s:
                    size_types.append("小体型")
                elif "中" in s:
                    size_types.append("中体型")
                elif "大" in s:
                    size_types.append("大体型")
    if not size_types:
        raise ConfigDriftError(
            "SIZE_TYPES 为空 — Excel 未解析到体型列 "
            "(地图遭遇表头或流派面板 R43-R45), 不允许硬编码三体型兜底"
        )

    race_types = list(cfg.RACE_TYPES) if cfg.RACE_TYPES else []
    if not race_types:
        for r in range(cfg.RACE_START_ROW, cfg.RACE_END_ROW + 1):
            v = ws.cell(row=r, column=2).value
            if v:
                race_types.append(str(v).strip())
    if not race_types:
        raise ConfigDriftError(
            "RACE_TYPES 为空 — resolve_columns 与流派面板 R33-R42 均未读到种族名, "
            "不允许硬编码种族列表兜底"
        )

    data_start = cfg.SCENE_DATA_START
    data_end = cfg.COL["SCENE_DATA_END"]
    if data_end > ws.max_row:
        print(f"{cfg.PRINT_PREFIX_WARN} SCENE_DATA_END={data_end} > ws.max_row={ws.max_row}, "
              f"(预期 {len(cfg.LEVELS_5STEP)}等级×{len(cfg.SCENES)}场景="
              f"{data_end - data_start + 1}行), "
              f"将截断到 max_row. 请检查 xlsx 数据是否完整.")
        data_end = ws.max_row
    cfg.qprint(f"{cfg.PRINT_PREFIX_READ} 场景价值 (从 Layer 3)")

    config = defaultdict(list)
    na_locations: List[Tuple[int, str, int]] = []
    for r in range(data_start, data_end + 1):
        level = ws.cell(row=r, column=cfg.COL["地图价值_等级"]).value
        if level is None:
            continue
        level_str = str(level).strip()
        if level_str.upper().startswith("LV"):
            level_str = level_str[2:]
        try:
            level = int(float(level_str))
        except (ValueError, TypeError):
            continue

        stype = str(ws.cell(row=r, column=cfg.COL["地图价值_场景类型"]).value or "").strip()
        if not stype or "合计" in stype:
            continue
        if level not in cfg.LEVELS_5STEP:
            continue

        map_name = str(ws.cell(row=r, column=cfg.COL["地图价值_关联地图"]).value or "").strip()
        if map_name == "—":
            map_name = ""

        value_raw = ws.cell(row=r, column=cfg.COL["地图价值_期望业务价值"]).value
        if value_raw is None or (isinstance(value_raw, str) and value_raw.strip() == "#N/A"):
            na_locations.append((r, stype, level))
            continue
        value_weight = _safe_float(value_raw)
        pve_raw = ws.cell(row=r, column=cfg.COL["地图价值_PVE占比"]).value
        pvp_raw = ws.cell(row=r, column=cfg.COL["地图价值_PVP占比"]).value
        pve_na = pve_raw is None or (isinstance(pve_raw, str) and pve_raw.strip() == "#N/A")
        pvp_na = pvp_raw is None or (isinstance(pvp_raw, str) and pvp_raw.strip() == "#N/A")
        if pve_na and pvp_na:
            raise ConfigDriftError(
                f"场景覆盖 Sheet R{r} (Lv{level} {stype}) PVE% 与 PVP% 双空, "
                f"业务真值应写入 PVE% / PVP% 列, 不允许默认纯 PVE 兜底"
            )
        if pve_na:
            pvp_ratio = _safe_float(pvp_raw)
            pve_ratio = 1.0 - pvp_ratio
        elif pvp_na:
            pve_ratio = _safe_float(pve_raw)
            pvp_ratio = 1.0 - pve_ratio
        else:
            pve_ratio = _safe_float(pve_raw)
            pvp_ratio = _safe_float(pvp_raw)

        config[level].append({
            "scene_type": stype, "map_name": map_name,
            "value_weight": value_weight,
            "pve_ratio": pve_ratio, "pvp_ratio": pvp_ratio,
        })

    levels = sorted(config.keys())
    total = sum(len(v) for v in config.values())
    if na_locations:
        value_col = cfg.COL["地图价值_期望业务价值"]
        value_letter = get_column_letter(value_col)
        loc_str = ", ".join(f"{value_letter}{r}({st},Lv{lv})" for r, st, lv in na_locations)
        raise ConfigDriftError(
            f"场景覆盖 Sheet {value_letter}{na_locations[0][0]}「期望业务价值」(col {value_col}) "
            f"为 None 或 \"#N/A\", {len(na_locations)}/{total} 处缺失: {loc_str}. "
            f"业务真值应直接写入 {value_letter} 列, 不允许手算兜底."
        )
    cfg.qprint(f"{cfg.PRINT_PREFIX_READ} 场景覆盖 Layer3 ({total} 行, 无 NA)")
    return levels, dict(config), attr_types, size_types, race_types


def _validate_layer_chain(scene_config, map_encounters, strict: bool = False):
    """校验 Layer 3 → Layer 2 数据链路完整性。返回 (missing_maps, unused_maps)。"""
    l3_refs = {}
    for level, scenes in scene_config.items():
        for s in scenes:
            mn = s.get("map_name", "")
            if mn and mn not in ("—", "", "None"):
                if mn not in l3_refs:
                    l3_refs[mn] = []
                l3_refs[mn].append((level, s.get("scene_type", "?")))

    l2_maps = set(map_encounters.keys())
    missing = {m: refs for m, refs in l3_refs.items() if m not in l2_maps}
    unused = [m for m in sorted(l2_maps) if m not in l3_refs]

    if missing:
        print(f"{cfg.PRINT_PREFIX_CHAIN_BLOCK} Layer 3 引用了 Layer 2 中不存在的地图:")
        for m, refs in sorted(missing.items()):
            locations = ", ".join(f"Lv{lv}/{sc}" for lv, sc in refs[:3])
            if len(refs) > 3:
                locations += f" ...共{len(refs)}处"
            print(f"  {cfg.PRINT_PREFIX_CHAIN_BREAK} \"{m}\" <- {locations}")
        print(f"  共 {len(missing)} 个断链, 禁止均匀分布兜底")
        raise RuntimeError(
            f"Layer3→Layer2 链路断链 {len(missing)} 处, "
            f"缺失地图示例: {list(sorted(missing.keys()))[:3]}"
        )
    else:
        cfg.qprint(f"{cfg.PRINT_PREFIX_CHAIN_OK} Layer 3 引用的 {len(l3_refs)} 个地图在 Layer 2 中全部存在")

    if unused:
        cfg.qprint(f"{cfg.PRINT_PREFIX_CHAIN_WARN} Layer 2 有 {len(unused)} 个地图未被 Layer 3 引用 (僵尸数据)")
    else:
        cfg.qprint(f"{cfg.PRINT_PREFIX_CHAIN_DONE} Layer 2 全部 {len(l2_maps)} 个地图均被 Layer 3 引用, 零冗余")

    return missing, unused


def aggregate_scene_coverage(scene_config, map_encounters, build_panel,
                             excel_pve_pvp=None):
    """Layer 3 × Layer 2 → PVE 覆盖输入。直接用 value_weight, 无场景价值调整系数。"""
    pve_config = defaultdict(list)
    per_level_weights = {}

    for level, scenes in scene_config.items():
        pve_total = 0.0
        pvp_total = 0.0

        for s in scenes:
            map_name = s["map_name"]
            scene_type = s["scene_type"]
            value_weight = s["value_weight"]

            pve_w = value_weight * s["pve_ratio"]
            pvp_w = value_weight * s["pvp_ratio"]
            pve_total += pve_w
            pvp_total += pvp_w

            if map_name and map_name in map_encounters and pve_w > cfg.PROGRAM_EPS:
                encounter = map_encounters[map_name]
                pve_config[level].append({
                    "name": f"{s['scene_type']}-{map_name}",
                    "type": s["scene_type"],
                    "level": level,
                    "weight": pve_w,
                    "attr": encounter["attr"].copy(),
                    "race": encounter["race"].copy(),
                    "size": encounter["size"].copy(),
                })
            elif pve_w > cfg.PROGRAM_EPS:
                raise ValueError(
                    f"{cfg.PRINT_PREFIX_WARN} Lv{level} {scene_type} → \"{map_name}\" "
                    f"在 Layer 2 中不存在, "
                    f"业务真值应写入地图遭遇 Sheet (Layer 2), 不允许均匀分布兜底"
                )

        total = pve_total + pvp_total
        if total > 0:
            per_level_weights[level] = {
                "pve_weight": pve_total / total,
                "pvp_weight": pvp_total / total,
            }
        else:
            raise ValueError(
                f"{cfg.PRINT_PREFIX_WARN} Lv{level} PVE+PVP 总权重={total}≤0, "
                f"业务数据缺失而非均衡, 不允许默认权重兜底"
            )

    # β̂（场景行推导口径）在 excel 等级曲线覆盖前留存，供体检对照双源漂移（2026-07-30）
    per_level_weights_hat = {lv: dict(w) for lv, w in per_level_weights.items()}

    if excel_pve_pvp:
        for lv, ratios in excel_pve_pvp.items():
            if lv in per_level_weights:
                per_level_weights[lv] = {
                    "pve_weight": ratios["pve_weight"],
                    "pvp_weight": ratios["pvp_weight"],
                }

    levels = sorted(pve_config.keys())
    cfg.qprint(f"{cfg.PRINT_PREFIX_AGGREGATE} Layer3→Layer4 PVE 场景聚合")
    return dict(pve_config), per_level_weights, levels, per_level_weights_hat


def _derive_playway_weights_from_scenes(scene_config, env_filter=None):
    """从 Layer 3 scene_config 按场景类型聚合总价值权重(已归一化)。"""
    totals = defaultdict(float)
    for _, scenes in scene_config.items():
        for s in scenes:
            stype = s.get("scene_type", "")
            val = s.get("value_weight", 0)
            if env_filter == "PVE":
                val *= s.get("pve_ratio", 1)
            elif env_filter == "PVP":
                val *= s.get("pvp_ratio", 0)
            mapped = cfg.SCENE_TO_PLAYWAY.get(stype, stype)
            totals[mapped] += val
    total = sum(totals.values())
    if total > 0:
        return {k: v / total for k, v in totals.items()}
    return {}


def read_pve_playway_weights(wb_or_config):
    """读取 PVE 玩法权重。支持从 scene_config dict 直接聚合。"""
    if isinstance(wb_or_config, dict):
        result = _derive_playway_weights_from_scenes(wb_or_config, "PVE")
        cfg.qprint(f"{cfg.PRINT_PREFIX_READ} PVE 玩法权重")
        return result
    cfg.qprint(f"{cfg.PRINT_PREFIX_READ} PVE 玩法权重 (playway 数据为空)")
    return {}


def read_pvp_playway_weights(wb_or_config):
    """读取 PVP 玩法权重。支持从 scene_config dict 直接聚合。"""
    if isinstance(wb_or_config, dict):
        result = _derive_playway_weights_from_scenes(wb_or_config, "PVP")
        cfg.qprint(f"{cfg.PRINT_PREFIX_READ} PVP 玩法权重")
        return result
    return {}
