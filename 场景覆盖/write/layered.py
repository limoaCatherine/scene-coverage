# -*- coding: utf-8 -*-
"""覆盖率结果：工具只写块标题与数据，不写说明句。"""
from __future__ import annotations

from pathlib import Path

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from 场景覆盖 import config as cfg

FONT = "微软雅黑"
NUM = "0.000000;-0.000000;0.000000"
PCT = "0.00%;-0.00%;0.00%"
FILL_L0 = PatternFill("solid", fgColor="F2F2F2")
FILL_L1 = PatternFill("solid", fgColor="E7E6E6")
FILL_L3 = PatternFill("solid", fgColor="D9D9D9")
FONT_L0 = Font(name=FONT, size=14, bold=True, color="1F4E79")
FONT_L1 = Font(name=FONT, size=12, bold=True, color="1F4E79")
FONT_L3 = Font(name=FONT, size=10, bold=True, color="404040")
FONT_TXT = Font(name=FONT, size=10, color="000000")
AL = Alignment(horizontal="left", vertical="center")
METRICS = ("输出乘区", "承伤乘区", "平衡指数")


def _dest(framework_path: str) -> str:
    path = str(framework_path)
    print(f"[闸] 覆盖率结果写回: {path}")
    return path


def _paint_l1(ws, r, c0, c1, title):
    for c in range(c0, c1 + 1):
        ws.cell(r, c).fill = FILL_L1
        ws.cell(r, c).alignment = AL
    ws.cell(r, c0, title).font = FONT_L1


def _hdr(ws, r, c, text):
    cell = ws.cell(r, c, text)
    cell.font = FONT_L3
    cell.fill = FILL_L3
    cell.alignment = AL


def _txt(ws, r, c, v):
    cell = ws.cell(r, c, v)
    cell.font = FONT_TXT
    cell.alignment = AL


def _num(ws, r, c, v, fmt=NUM):
    cell = ws.cell(r, c, None if v is None else float(v))
    cell.font = FONT_TXT
    cell.alignment = AL
    cell.number_format = fmt


def _trio_headers(ws, r_name, r_metric, c0, builds):
    for i, b in enumerate(builds):
        base = c0 + i * 3
        _hdr(ws, r_name, base, b)
        ws.cell(r_name, base + 1).fill = FILL_L3
        ws.cell(r_name, base + 2).fill = FILL_L3
        if base + 2 > base:
            try:
                ws.merge_cells(start_row=r_name, start_column=base, end_row=r_name, end_column=base + 2)
            except ValueError:
                pass
        for j, m in enumerate(METRICS):
            _hdr(ws, r_metric, base + j, m)


def _trio_values(ws, r, c0, rec):
    rec = rec or {}
    _num(ws, r, c0, rec.get("输出乘区"))
    _num(ws, r, c0 + 1, rec.get("承伤乘区"))
    _num(ws, r, c0 + 2, rec.get("平衡指数"))


def write_layered(framework_path: str, rankings: dict, layered: dict, meta_note: str = "") -> str:
    del rankings, meta_note
    path = _dest(framework_path)
    import openpyxl

    wb = openpyxl.load_workbook(path)
    ws = wb[cfg.SHEET_OUTPUT] if cfg.SHEET_OUTPUT in wb.sheetnames else wb.create_sheet(cfg.SHEET_OUTPUT)
    for mr in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(mr))
    if ws.max_row >= 1:
        ws.delete_rows(1, ws.max_row)

    builds = list(layered["builds"])
    n = len(builds)
    by_mode = layered["by_mode"]
    by_lv = layered["by_level"]
    by_play = layered["by_playway"]
    by_map = layered.get("by_map") or {}
    by_elem = layered.get("by_elem") or {}
    by_size = layered.get("by_size") or {}
    map_meta = layered.get("map_meta") or {}
    race_env = layered.get("race_env") or {}
    size_env = layered.get("size_env") or {}
    race_dist = layered.get("race_dist") or {}
    size_dist = layered.get("size_dist") or {}
    race_types = layered.get("race_types") or []
    size_types = layered.get("size_types") or []
    attr_types = layered.get("attr_types") or []
    plays = layered["play_names"]
    maps = list(map_meta.keys()) or list(by_map.keys())

    w_build = n * 3
    c_board, w_board = 1, 5
    c_lv = c_board + w_board + 1
    w_lv = 1 + w_build
    c_play = c_lv + w_lv + 1
    w_play = 1 + w_build
    c_mode = c_play + w_play + 1
    w_mode = 1 + w_build
    c_map = c_mode + w_mode + 1
    w_map = 3 + w_build
    c_race = c_map + w_map + 1
    w_race = 2 + n
    c_size = c_race + w_race + 1
    w_size = 2 + w_build
    c_elem = c_size + w_size + 1
    w_elem = 1 + w_build
    used = c_elem + w_elem - 1

    blocks = [
        (c_board, c_board + w_board - 1, "看板·综合"),
        (c_lv, c_lv + w_lv - 1, "分等级"),
        (c_play, c_play + w_play - 1, "分玩法"),
        (c_mode, c_mode + w_mode - 1, "分模式"),
        (c_map, c_map + w_map - 1, "分地图"),
        (c_race, c_race + w_race - 1, "三维·种族"),
        (c_size, c_size + w_size - 1, "三维·体型"),
        (c_elem, c_elem + w_elem - 1, "三维·元素"),
    ]

    ws.cell(1, 1, "覆盖率结果").font = FONT_L0
    ws["A1"].fill = FILL_L0
    ws["A1"].alignment = AL
    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 8

    for c0, c1, title in blocks:
        _paint_l1(ws, 3, c0, c1, title)

    for i, h in enumerate(("流派", "输出乘区", "承伤乘区", "平衡指数", "价值权重")):
        _hdr(ws, 4, c_board + i, h)
        _hdr(ws, 5, c_board + i, "")
    mode_pve = by_mode.get("PVE", {})
    mode_pvp = by_mode.get("PVP", {})
    for i, b in enumerate(builds):
        r = 6 + i
        a, q = mode_pve.get(b, {}), mode_pvp.get(b, {})
        _txt(ws, r, c_board, b)
        _num(ws, r, c_board + 1, _avg(a.get("输出乘区"), q.get("输出乘区")))
        _num(ws, r, c_board + 2, _avg(a.get("承伤乘区"), q.get("承伤乘区")))
        _num(ws, r, c_board + 3, _avg(a.get("平衡指数"), q.get("平衡指数")))
        _num(ws, r, c_board + 4, (a.get("价值权重") or 0) + (q.get("价值权重") or 0))

    def _keys_block(c0, key_h, keys, store, extra_hdr=(), extra_fn=None, key_fmt=None):
        _hdr(ws, 4, c0, key_h)
        _hdr(ws, 5, c0, key_h)
        for j, h in enumerate(extra_hdr):
            _hdr(ws, 4, c0 + 1 + j, h)
            _hdr(ws, 5, c0 + 1 + j, h)
        off = 1 + len(extra_hdr)
        _trio_headers(ws, 4, 5, c0 + off, builds)
        for i, key in enumerate(keys):
            rr = 6 + i
            _txt(ws, rr, c0, key)
            if key_fmt:
                ws.cell(rr, c0).number_format = key_fmt
            if extra_fn:
                for j, val in enumerate(extra_fn(key)):
                    if isinstance(val, (int, float)) and not isinstance(val, bool):
                        _num(ws, rr, c0 + 1 + j, val)
                    else:
                        _txt(ws, rr, c0 + 1 + j, val)
            row = store.get(key, {})
            for j, b in enumerate(builds):
                _trio_values(ws, rr, c0 + off + j * 3, row.get(b))

    _keys_block(c_lv, "等级", list(range(1, 61)), by_lv, key_fmt="0")
    _keys_block(c_play, "玩法", plays, by_play)
    _keys_block(c_mode, "模式", ("PVE", "PVP"), by_mode)

    def _map_extra(name):
        meta = map_meta.get(name, {})
        return (meta.get("玩法", ""), meta.get("等级", ""))

    _keys_block(c_map, "地图", maps, by_map, extra_hdr=("玩法", "等级"), extra_fn=_map_extra)
    for i, name in enumerate(maps):
        lv = map_meta.get(name, {}).get("等级")
        if isinstance(lv, (int, float)):
            ws.cell(6 + i, c_map + 2).number_format = "0"

    _hdr(ws, 4, c_race, "种族")
    _hdr(ws, 5, c_race, "种族")
    _hdr(ws, 4, c_race + 1, "环境占比")
    _hdr(ws, 5, c_race + 1, "环境占比")
    for j, b in enumerate(builds):
        _hdr(ws, 4, c_race + 2 + j, b)
        _hdr(ws, 5, c_race + 2 + j, "三维权重")
    for i, rname in enumerate(race_types):
        rr = 6 + i
        _txt(ws, rr, c_race, rname)
        _num(ws, rr, c_race + 1, race_env.get(rname, 0.0), PCT)
        for j, b in enumerate(builds):
            _num(ws, rr, c_race + 2 + j, (race_dist.get(b) or {}).get(rname, 0.0), PCT)

    _hdr(ws, 4, c_size, "体型")
    _hdr(ws, 5, c_size, "体型")
    _hdr(ws, 4, c_size + 1, "环境占比")
    _hdr(ws, 5, c_size + 1, "环境占比")
    if by_size:
        _trio_headers(ws, 4, 5, c_size + 2, builds)
    else:
        for j, b in enumerate(builds):
            _hdr(ws, 4, c_size + 2 + j, b)
            _hdr(ws, 5, c_size + 2 + j, "三维权重")
    for i, sname in enumerate(size_types):
        rr = 6 + i
        _txt(ws, rr, c_size, sname)
        _num(ws, rr, c_size + 1, size_env.get(sname, 0.0), PCT)
        row = by_size.get(sname, {})
        for j, b in enumerate(builds):
            rec = row.get(b)
            if rec and rec.get("输出乘区") is not None:
                _trio_values(ws, rr, c_size + 2 + j * 3, rec)
            else:
                _num(ws, rr, c_size + 2 + j, (size_dist.get(b) or {}).get(sname, 0.0), PCT)

    _keys_block(c_elem, "元素", attr_types, by_elem)

    for c in range(1, used + 1):
        ws.column_dimensions[get_column_letter(c)].width = 11
    ws.column_dimensions[get_column_letter(c_map)].width = 16
    ws.sheet_view.showGridLines = False
    wb.active = ws
    wb.save(path)
    wb.close()
    return path


def _avg(a, b):
    vals = [v for v in (a, b) if v is not None]
    if not vals:
        return None
    return sum(vals) / len(vals)
