
def _publish_out(name: str):
    from pathlib import Path
    dest = Path(__file__).resolve().parents[2] / "out" / name
    dest.parent.mkdir(parents=True, exist_ok=True)
    return dest

# results.py — 只写「覆盖率结果」Sheet
from __future__ import annotations

from typing import List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from 场景覆盖 import config as cfg

_THIN = Side(style="thin", color="B0B0B0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _rel_fill(value: float, dim: str = "流派") -> Optional[PatternFill]:
    """指数 (~1.0) → PASS/WARN/FAIL 填色。"""
    th = (cfg.RANK_REL_THRESHOLDS or {}).get(dim)
    if not th:
        return None
    lo, hi = th["pass"]
    if lo <= value <= hi:
        return cfg.FILL_PASS
    wlo, whi = th["warn"]
    if wlo <= value <= whi:
        return cfg.FILL_WARN
    return cfg.FILL_FAIL


def _fill_row_rel(value: float, row_vals: List[float], dim: str, invert: bool = False) -> PatternFill:
    """行内相对均值填色：rel = v / 行均值（1.0 = 该行维内均值）。

    原始效用值均值不在 1.0（如武器 util∈[0.5,1.0]，均值≈0.85），
    直接套「1.0=均值」阈值带会系统性错色——数据块与排名块互相矛盾（2026-07-30 修）。
    invert=True 用于守方/挨打量（高=坏，红绿对调，与排名 S=(攻/均攻)×(均守/守) 同向）。
    """
    m = sum(row_vals) / len(row_vals) if row_vals else 0.0
    rel = float(value) / m if m > cfg.PROGRAM_EPS else 1.0
    if invert:
        rel = 2.0 - rel
    return _rel_fill(rel, dim)


def _share_fill(value: float) -> Optional[PatternFill]:
    """场景占比热力；过小不着色。"""
    if value <= 0:
        return None
    for upper, fill in cfg.SHARE_FILL_BANDS:
        if value < upper:
            return fill
    return cfg.FILL_SHARE_4


def _apply_border(ws, r1: int, c1: int, r2: int, c2: int) -> None:
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(row=r, column=c).border = _BORDER


def write_results(
    framework_path: str,
    detailed: dict,
    pure_dist: dict,
    quad: dict,
    bfi: dict,
    size_weapon_util: dict,
    size_types: List[str],
    rankings: dict,
    meta_note: str = "",
    size_hit_util: dict | None = None,
):
    """写 5 维排名 (R1+) + 数据块 (R25+)。正式簿只读；结果落到沙盒副本。"""
    from pathlib import Path
    import shutil

    src = Path(framework_path)
    dest = _publish_out("sandbox-覆盖率结果.xlsx")
    dest.parent.mkdir(parents=True, exist_ok=True)
    if src.resolve() != dest.resolve():
        shutil.copy2(src, dest)
        framework_path = str(dest)
        print(f"[闸] 覆盖率结果只写沙盒: {dest}")
    wb = openpyxl.load_workbook(framework_path)
    ws = wb[cfg.SHEET_OUTPUT] if cfg.SHEET_OUTPUT in wb.sheetnames else wb.create_sheet(cfg.SHEET_OUTPUT)

    # 清掉旧合并（标题跨列），再全表清空。
    # 必须整表清：旧版式（三维评估/中文状态/排名@R24）与新代码块位完全不同，
    # 只清 R25+ 会让旧 R1-R24 的块标题与新块混排（新旧版式漂移残留，2026-07-30 登记）。
    for mr in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(mr))
    if ws.max_row >= 1:
        ws.delete_rows(1, ws.max_row)

    bold = Font(bold=True)
    title_font = Font(bold=True, size=11, color="FFFFFF", name="Calibri")
    header_font = Font(bold=True, size=10, color="FFFFFF", name="Calibri")
    num_fmt = cfg.NUM_FMT
    int_fmt = cfg.INT_FMT

    levels = sorted(pure_dist.keys())
    if not levels:
        raise ValueError("无有效等级数据")
    attr_types = list(detailed["PVE"]["attr_attacker"][levels[0]].keys())
    weapon_types = list(size_weapon_util["PVE"][levels[0]].keys())

    # (r1, c1, r2, c2) 供最后统一画线框
    blocks: List[Tuple[int, int, int, int]] = []

    def _title(r, c, text, n_cols: int, *, fill=None, font=None):
        """标题跨块宽合并 + 蓝底白字（可覆写样式）。"""
        end = c + n_cols - 1
        use_fill = fill or cfg.FILL_TITLE
        use_font = font or title_font
        cell = ws.cell(row=r, column=c, value=text)
        cell.font = use_font
        cell.fill = use_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")
        for ci in range(c + 1, end + 1):
            sc = ws.cell(row=r, column=ci, value=None)
            sc.fill = use_fill
            sc.font = use_font
        if end > c:
            ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=end)

    def _section_banner(r, text: str, n_cols: int = 32):
        """明细区组间分隔条：深蓝底 + 白字，强化专业分块观感。"""
        banner_font = Font(bold=True, size=11, color="FFFFFF")
        banner_fill = PatternFill("solid", fgColor="2F5496")
        _title(r, 1, text, n_cols, fill=banner_fill, font=banner_font)
        ws.row_dimensions[r].height = 18

    def _hdr(r, c, headers):
        for i, h in enumerate(headers):
            cell = ws.cell(row=r, column=c + i, value=h)
            cell.font = header_font
            cell.fill = cfg.FILL_HEADER
            cell.alignment = Alignment(horizontal="center", vertical="center")

    def _label(r, c, value):
        cell = ws.cell(row=r, column=c, value=value)
        cell.font = bold
        cell.fill = cfg.FILL_LABEL
        if isinstance(value, int):
            cell.number_format = int_fmt
        return cell

    def _num(r, c, value, *, fill: Optional[PatternFill] = None, decimals: int = 4):
        cell = ws.cell(row=r, column=c, value=round(float(value), decimals))
        cell.number_format = num_fmt
        if fill is not None:
            cell.fill = fill
        return cell

    def _num_rel(r, c, value, dim: str = "流派"):
        # 填色以舍入后的显示值为准（原值 0.199996 显示 0.2000 却按 <0.20 带填色的边界错色，2026-07-30 修）
        v = round(float(value), 4)
        return _num(r, c, v, fill=_rel_fill(v, dim))

    def _num_share(r, c, value):
        v = round(float(value), 4)
        return _num(r, c, v, fill=_share_fill(v))

    # R1–R22: 五维排名看板；R25+: 明细数据块（无新鲜度横幅）
    DATA_TITLE_ROW = 25
    cur_col = 1
    GAP = 1

    def _new_block(title: str, n_data_cols: int, title_row: int, n_body_rows: int):
        """n_data_cols = 数据列数（不含首列标签）；块总宽 = 1+n_data_cols。"""
        nonlocal cur_col
        if cur_col > 1:
            cur_col += GAP
        start = cur_col
        width = n_data_cols + 1
        _title(title_row, start, title, width)
        # title + header + body
        r2 = title_row + 1 + n_body_rows
        blocks.append((title_row, start, r2, start + width - 1))
        cur_col = start + width
        return start

    def _level_rows(block_col, start_row):
        for ri, lv in enumerate(levels):
            r = start_row + ri
            _label(r, block_col, int(lv))
            yield r, lv

    cross = bfi.get("跨等级综合") or {}
    bfi_builds = sorted(cross.keys(), key=lambda b: -float(cross[b].get("bfi", 0)))
    if not bfi_builds:
        raise ValueError(f"{cfg.PRINT_PREFIX_WARN} 跨等级综合 BFI 为空")

    def _build_rows(block_col, start_row):
        for ri, build in enumerate(bfi_builds):
            r = start_row + ri
            _label(r, block_col, build)
            yield r, build

    n_lv = len(levels)
    n_bfi = len(bfi_builds)
    h1 = max(n_lv, n_bfi) + 2
    h2 = h3 = n_lv + 2
    group_rows = [DATA_TITLE_ROW]
    for h in (h1, h2, h3):
        group_rows.append(group_rows[-1] + h + 2)
    # 供 verify 使用
    cfg.BLOCK_GROUP_ROWS = list(group_rows)

    # ---- 组1: 场景属性乘区 + 流派平衡指数 ----
    g1 = group_rows[0]
    cur_col = 1
    B2 = _new_block(cfg.BLOCK_TITLE_QUAD, 6, g1, n_lv)
    _hdr(g1 + 1, B2, list(cfg.BLOCK_HDR_QUAD))
    inds = [("PVE", "attacker"), ("PVE", "defender"), ("PVP", "attacker"),
            ("PVP", "defender"), ("综合", "attacker"), ("综合", "defender")]
    for r, lv in _level_rows(B2, g1 + 2):
        for ci, (env, role) in enumerate(inds):
            v = round(float(quad[env][role].get(lv, 0)), 4)
            if role == "defender":
                # 承伤高=坏，红绿对调
                _num(r, B2 + 1 + ci, v, fill=_rel_fill(2.0 - v))
            else:
                _num_rel(r, B2 + 1 + ci, v)

    B3 = _new_block(cfg.BLOCK_TITLE_BFI_CROSS, 3, g1, n_bfi)
    _hdr(g1 + 1, B3, list(cfg.BLOCK_HDR_BFI))
    for r, build in _build_rows(B3, g1 + 2):
        d = cross[build]
        dv = round(float(d["defender"]), 4)
        _num_rel(r, B3 + 1, d["attacker"])
        _num(r, B3 + 2, dv, fill=_rel_fill(2.0 - dv))
        _num_rel(r, B3 + 3, d["bfi"])

    for env in ("综合", "PVE", "PVP"):
        Bi = _new_block(cfg.BLOCK_TITLE_BFI_BY_LV.format(env=env), n_lv, g1, n_bfi)
        _hdr(g1 + 1, Bi, ["流派"] + [int(lv) for lv in levels])
        for r, build in _build_rows(Bi, g1 + 2):
            for ci, lv in enumerate(levels):
                d = bfi.get(env, {}).get(lv, {}).get(build, {})
                _num_rel(r, Bi + 1 + ci, d.get("bfi", 0))

    # ---- 组2: 场景占比 ----
    cur_col = 1
    g2 = group_rows[1]
    for dim, n_cols in (("属性", 10), ("体型", 3), ("种族", 10)):
        for env in ("综合", "PVE", "PVP"):
            Bi = _new_block(
                cfg.BLOCK_TITLE_SCENE_DIST.format(env=env, dim=dim),
                n_cols, g2, n_lv,
            )
            headers = list(pure_dist[levels[0]][env][dim].keys())[:n_cols]
            _hdr(g2 + 1, Bi, ["等级"] + headers)
            for r, lv in _level_rows(Bi, g2 + 2):
                dist = pure_dist[lv][env][dim]
                for ci, h in enumerate(headers):
                    _num_share(r, Bi + 1 + ci, dist.get(h, 0))

    # ---- 组3: 属性输出/承伤乘区 ----
    cur_col = 1
    g3 = group_rows[2]
    for role, key, title_tpl in (
        ("攻方", "attr_attacker", cfg.BLOCK_TITLE_ATTR_ATK),
        ("守方", "attr_defender", cfg.BLOCK_TITLE_ATTR_DEF),
    ):
        for env in ("综合", "PVE", "PVP"):
            Bi = _new_block(
                title_tpl.format(env=env), len(attr_types), g3, n_lv
            )
            _hdr(g3 + 1, Bi, ["等级"] + attr_types)
            for r, lv in _level_rows(Bi, g3 + 2):
                d = detailed[env][key][lv]
                row_vals = [round(float(d.get(a, 0)), 4) for a in attr_types]
                for ci, a in enumerate(attr_types):
                    v = row_vals[ci]
                    _num(r, Bi + 1 + ci, v,
                         fill=_fill_row_rel(v, row_vals, "属性", invert=(key == "attr_defender")))

    # ---- 组4: 武器输出乘区 + 体型承伤乘区 ----
    cur_col = 1
    g4 = group_rows[3]
    for i, env in enumerate(("综合", "PVE", "PVP")):
        if i > 0:
            cur_col += 2
        Bi = _new_block(
            cfg.BLOCK_TITLE_WEAPON_ATK.format(env=env),
            len(weapon_types), g4, n_lv,
        )
        _hdr(g4 + 1, Bi, ["等级"] + weapon_types)
        for r, lv in _level_rows(Bi, g4 + 2):
            d = size_weapon_util.get(env, {}).get(lv, {})
            row_vals = [round(float(d.get(w, 0)), 4) for w in weapon_types]
            for ci, w in enumerate(weapon_types):
                v = row_vals[ci]
                _num(r, Bi + 1 + ci, v, fill=_fill_row_rel(v, row_vals, "武器"))

    if size_hit_util:
        cur_col += 2
        for env in ("综合", "PVE", "PVP"):
            Bi = _new_block(
                cfg.BLOCK_TITLE_SIZE_HIT.format(env=env),
                len(size_types), g4, n_lv,
            )
            _hdr(g4 + 1, Bi, ["等级"] + list(size_types))
            for r, lv in _level_rows(Bi, g4 + 2):
                d = size_hit_util.get(env, {}).get(lv, {})
                row_vals = [round(float(d.get(sz, 0)), 4) for sz in size_types]
                for ci, sz in enumerate(size_types):
                    v = row_vals[ci]
                    # 承伤高=坏，红绿对调
                    _num(r, Bi + 1 + ci, v,
                         fill=_fill_row_rel(v, row_vals, "体型", invert=True))

    for c in range(1, max(cur_col + 20, 40)):
        ws.column_dimensions[get_column_letter(c)].width = 11
    ws.column_dimensions["A"].width = 14

    # ---- 顶部五维看板 ----
    notes = []
    if meta_note or rankings.get("meta_note"):
        notes.append(meta_note or rankings.get("meta_note"))
    if rankings.get("race_mode") == "coverage_share":
        notes.append("种族块=场景出现占比÷十族均值(无种族克制矩阵，非战斗强弱)")
    notes.extend(cfg.RESULT_NOTES_BASE)
    notes.append("数据块填色=行内相对均值(1.0=该行平均)；承伤列红绿相反")
    notes.append("看板=五维排名(R1)；明细自 R25 起")
    note_col = 34
    if notes:
        nc = ws.cell(row=1, column=note_col, value="[口径] " + "；".join(notes))
        nc.font = Font(italic=True, size=9, color="808080")

    def _pct_cell(r, c, pct, fill):
        cell = ws.cell(row=r, column=c, value=f"{float(pct):+.1f}%")
        cell.number_format = "@"
        cell.alignment = Alignment(horizontal="center")
        if fill is not None:
            cell.fill = fill
        return cell

    def _status_cell(r, c, status, fill):
        cell = ws.cell(row=r, column=c, value=status)
        cell.alignment = Alignment(horizontal="center")
        cell.font = bold
        if fill is not None:
            cell.fill = fill
        return cell

    def _write_rank_row(r, col_start, block_id, row, race_mode: str):
        """按块类型写一行双轨数据；返回状态列偏移。"""
        name = row.get("name", "")
        fill = row.get("fill")
        _label(r, col_start, name)
        if not name:
            return 0
        if block_id in ("build", "elem") or (
            block_id == "race" and race_mode == "matrix"
        ):
            _num(r, col_start + 1, row["atk"])
            _num(r, col_start + 2, row["def"])
            _num(r, col_start + 3, row["rel"], fill=fill)
            _pct_cell(r, col_start + 4, row["vs_pct"], fill)
            _status_cell(r, col_start + 5, row["status"], fill)
            return 5
        if block_id == "size":
            # 承伤乘区 | 相对均值差 | 抗性指数 | 判定
            _num(r, col_start + 1, row["real"])
            _num(r, col_start + 2, row["vs_pct"], fill=fill)
            _num(r, col_start + 3, row["rel"], fill=fill)
            _status_cell(r, col_start + 4, row["status"], fill)
            return 4
        # weapon / race: 乘区或占比 | 指数 | 偏离均值% | 判定
        _num(r, col_start + 1, row["real"])
        _num(r, col_start + 2, row["rel"], fill=fill)
        _pct_cell(r, col_start + 3, row["vs_pct"], fill)
        _status_cell(r, col_start + 4, row["status"], fill)
        return 4

    summary = []
    race_mode = rankings.get("race_mode") or "coverage_share"
    for block_id, title, col_start, n_data in cfg.RANK_BLOCK_LAYOUT:
        pairs = list(rankings["blocks"].get(block_id, []))
        if len(pairs) > n_data:
            pairs = pairs[:n_data]
        while len(pairs) < n_data:
            pairs.append({"name": ""})

        if block_id == "race" and race_mode == "matrix":
            headers = list(cfg.RANK_HEADERS_RACE_MATRIX)
        else:
            headers = list(cfg.RANK_BLOCK_HEADERS[block_id])
        width = len(headers)
        _title(1, col_start, title, width)
        _hdr(2, col_start, headers)
        last_data_r = 2
        for ri, row in enumerate(pairs):
            r = 3 + ri
            if not isinstance(row, dict):
                # 旧 tuple 形态不应再出现
                raise TypeError(f"排名行须为 dict, 收到 {type(row)}")
            _write_rank_row(r, col_start, block_id, row, race_mode)
            if row.get("name"):
                last_data_r = r

        named = [p for p in pairs if p.get("name")]
        blocks.append((1, col_start, last_data_r, col_start + width - 1))
        summary.append(f"{title}={len(named)}条")

    for r1, c1, r2, c2 in blocks:
        _apply_border(ws, r1, c1, r2, c2)

    wb.save(framework_path)
    wb.close()
    cfg.qprint(
        f"{cfg.PRINT_PREFIX_OUTPUT} 5 ranking + 数据块 → '{cfg.SHEET_OUTPUT}' "
        f"({', '.join(summary)})"
    )
    return framework_path
